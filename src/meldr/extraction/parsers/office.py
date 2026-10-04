"""PDF (pypdf), DOCX (python-docx), and XLSX (openpyxl) parsers — extra `docs`."""

from __future__ import annotations

import io

from meldr.core.errors import MissingExtraError
from meldr.extraction.parsers import DocElement


def parse_pdf(data: bytes) -> list[DocElement]:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise MissingExtraError("docs", "Parsing PDF files") from e
    reader = PdfReader(io.BytesIO(data))
    out = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append(DocElement(kind="paragraph", text=text, page=i))
    return out


def parse_docx(data: bytes) -> list[DocElement]:
    try:
        import docx
    except ImportError as e:
        raise MissingExtraError("docs", "Parsing DOCX files") from e
    document = docx.Document(io.BytesIO(data))
    out: list[DocElement] = []
    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        if style.lower().startswith(("heading", "title")):
            out.append(DocElement(kind="heading", text=text))
        else:
            out.append(DocElement(kind="paragraph", text=text))
    for table in document.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if rows:
            out.append(DocElement(kind="table", text="\n".join(rows)))
    return out


def parse_xlsx(data: bytes) -> list[DocElement]:
    try:
        import openpyxl
    except ImportError as e:
        raise MissingExtraError("docs", "Parsing XLSX files") from e
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        lines = [
            " | ".join("" if v is None else str(v) for v in row)
            for row in ws.iter_rows(values_only=True)
            if any(v is not None for v in row)
        ]
        if lines:
            out.append(DocElement(kind="table", text="\n".join(lines), sheet=ws.title))
    return out
