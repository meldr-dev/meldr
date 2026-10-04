from __future__ import annotations

import json
import shutil
from pathlib import Path

import polars as pl
import pytest
from typer.testing import CliRunner

from meldr.cli.main import app
from meldr.core.canonical import SchemaRegistry
from meldr.core.errors import ExternalToolError
from meldr.eval.runner import builtin_factories, check_baseline, run_eval
from meldr.eval.synth import make_variants
from meldr.extraction import media
from meldr.extraction.chunker import chunk
from meldr.extraction.parsers import DocElement, detect_kind, parse
from meldr.extraction.transcribe import TranscriptSegment
from meldr.llm.gateway import LLMGateway
from meldr.llm.providers.fake import FakeProvider
from meldr.pipeline.oracle import heuristic_oracle

runner = CliRunner()


def test_cli_flow(tmp_path: Path, erp_df: pl.DataFrame) -> None:
    cfg = tmp_path / "meldr.yaml"
    cfg.write_text(
        f"storage: {{root: {tmp_path / 'state'}}}\nllm: {{providers: {{fake: {{type: fake}}}}}}\n"
    )
    data = tmp_path / "erp.csv"
    erp_df.write_csv(data)
    c = ["--config", str(cfg)]
    assert "meldr" in runner.invoke(app, ["--version"]).output
    assert runner.invoke(app, ["profile", "--file", str(data), *c]).exit_code == 0
    run = runner.invoke(app, ["run", "-t", "Organization", "--file", str(data), *c])
    assert run.exit_code == 0 and "erp → Organization" in run.output
    rev = runner.invoke(
        app, ["review", "erp.Organization", "--approve", "org_type", "--user", "ana", *c]
    )
    assert rev.exit_code == 0 and "v2" in rev.output
    assert "erp.Organization" in runner.invoke(app, ["specs", "list", *c]).output
    diff = runner.invoke(app, ["specs", "diff", "erp.Organization", "--from", "1", "--to", "2", *c])
    assert '"org_type"' in diff.output
    bad = runner.invoke(app, ["review", "nope.Organization", *c])
    assert bad.exit_code == 1


def test_eval_and_baseline(tmp_path: Path, registry: SchemaRegistry) -> None:
    from meldr._async import run_sync

    org = registry.get("Organization")
    assert make_variants(org, 2, seed=1)[0].source.equals(make_variants(org, 2, seed=1)[0].source)
    report = run_sync(run_eval([org], builtin_factories(), variants=3, rows=40))
    by = {r.provider: r for r in report.results}
    assert by["fake"].f1 == 1.0 and by["heuristic"].auto_approved == 0
    baseline = tmp_path / "b.json"
    baseline.write_text(json.dumps({"heuristic": {"f1": 1.1}}))
    assert check_baseline(report, baseline)


def test_parsers_and_chunker(tmp_path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("Scope", level=1)
    d.add_paragraph("Build a clinic.")
    path = tmp_path / "a.docx"
    d.save(str(path))
    data = path.read_bytes()
    assert detect_kind("renamed.bin", data) == "docx"
    elements, _ = parse("a.docx", data)
    assert [e.kind for e in elements] == ["heading", "paragraph"]
    html, _ = parse("x.html", b"<html><script>x()</script><h1>T</h1><p>Body</p></html>")
    assert [e.text for e in html] == ["T", "Body"]
    segs = chunk(
        [DocElement("heading", "H"), DocElement("paragraph", "word " * 400, page=2)],
        record_id="r",
        source_id="s",
        target_tokens=100,
        overlap_tokens=10,
    )
    assert len(segs) > 2 and segs[0].title == "H" and segs[1].locator.page == 2


async def test_media_windows_and_captions(monkeypatch: pytest.MonkeyPatch) -> None:
    gw = LLMGateway(FakeProvider(mode="oracle", oracle=heuristic_oracle))
    caps = await media.caption_frames(gw, [media.Frame(ms=5000, jpeg=b"x")])
    assert caps[0].frame_ms == 5000
    segs = media.windows(
        [TranscriptSegment(start_ms=0, end_ms=70_000, text="hello")],
        caps,
        record_id="v",
        source_id="s",
        duration_ms=90_000,
        window_s=60,
    )
    assert [(s.locator.start_ms, s.locator.end_ms) for s in segs] == [(0, 60_000), (60_000, 90_000)]
    assert "Visual" in segs[0].text
    monkeypatch.setattr(shutil, "which", lambda name: None)
    with pytest.raises(ExternalToolError, match="ffmpeg"):
        media.probe(Path("x.mp4"))
